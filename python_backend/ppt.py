from __future__ import annotations

from pathlib import Path
from uuid import uuid4


def generate_ppt(query: str, output_dir: Path) -> Path:
    """Create a small, valid PPTX from a user's request.

    Generation errors must propagate to the API instead of producing an empty
    file that looks like a successful download. Long requests are wrapped and
    capped so the default template remains readable.
    """
    if not isinstance(query, str) or not query.strip():
        raise ValueError("PPT 需求不能为空")

    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / f"{uuid4().hex}.pptx"

    from pptx import Presentation
    from pptx.util import Pt

    request = " ".join(query.split())
    # Keep the first slide useful even when a pasted prompt is extremely long.
    request = request[:1200]
    presentation = Presentation()
    slides = [
        ("AI 智能演示", request),
        ("核心内容", "需求分析、方案设计与落地路径"),
        ("实施计划", "准备、执行、验证、迭代"),
    ]
    layout = presentation.slide_layouts[1]
    for title, body in slides:
        slide = presentation.slides.add_slide(layout)
        title_shape = slide.shapes.title
        if title_shape is None:
            raise RuntimeError("PPT 模板缺少标题占位符")
        title_shape.text = title
        body_shape = slide.placeholders[1]
        frame = body_shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = Pt(8)
        frame.margin_right = Pt(8)
        frame.margin_top = Pt(6)
        frame.margin_bottom = Pt(6)
        paragraph = frame.paragraphs[0]
        paragraph.text = body
        paragraph.font.size = Pt(20 if len(body) > 500 else 22)

    presentation.save(output)

    # Re-open the artifact to catch truncated or otherwise invalid files before
    # the API publishes a download URL.
    if not output.is_file() or output.stat().st_size == 0:
        raise RuntimeError("PPT 文件写出失败")
    try:
        Presentation(output)
    except Exception as exc:
        output.unlink(missing_ok=True)
        raise RuntimeError("生成的 PPT 文件无法打开") from exc
    return output
